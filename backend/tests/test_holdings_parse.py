"""Regression tests for the generic SEBI-template sheet parser
(app/ingestion/holdings/base.parse_standard_sheet).

Locks the Nippon-driven fixes: an internal code in column A must not be taken as
the scheme name, ISIN may sit before Name (columns found by label, not position),
'% to NAV' may be a fraction, and dates like 'June 30,2026' (no space after the
comma) must parse.

Run: ./venv/bin/python -m pytest tests/test_holdings_parse.py -v
"""
import datetime

import openpyxl

from app.ingestion.holdings import base


def _nippon_style_sheet():
    """Minimal workbook mimicking Nippon: code in col A, ISIN before Name, % as a
    fraction, date 'June 30,2026'."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["RLMF015", "NIPPON INDIA PHARMA FUND (An Open Ended Equity Scheme)"])
    ws.append([None, "Monthly Portfolio Statement as on June 30,2026"])
    ws.append([None])
    ws.append([None, "ISIN", "Name of the Instrument", "Industry / Rating",
               "Quantity", "Market/Fair Value ( Rs. in Lacs)", "% to NAV"])
    ws.append([None, "Equity & Equity related"])
    ws.append(["SPIL03", "INE044A01036", "Sun Pharmaceutical Industries Limited",
               "Pharmaceuticals & Biotechnology", 6206349, 115593.25, 0.1278])
    ws.append(["LUPL02", "INE326A01037", "Lupin Limited",
               "Pharmaceuticals & Biotechnology", 2783991, 67344.74, 0.0745])
    ws.append([None, "Sub Total", None, None, None, None, 0.2023])   # no ISIN -> skipped
    ws.append([None, "TREPS / Reverse Repo", None, None, None, None, 0.03])  # skipped
    return ws


def test_parses_nippon_style_layout():
    name, as_of, rows = base.parse_standard_sheet(_nippon_style_sheet())
    # scheme name is the fund, not the "RLMF015" code
    assert name.startswith("NIPPON INDIA PHARMA FUND")
    # date with no space after the comma still parses
    assert as_of == datetime.date(2026, 6, 30)
    # only the two ISIN-bearing rows are holdings; subtotal / TREPS skipped
    assert len(rows) == 2
    top = rows[0]
    assert top["instrument_name"] == "Sun Pharmaceutical Industries Limited"
    assert top["isin"] == "INE044A01036"
    assert top["sector"] == "Pharmaceuticals & Biotechnology"
    # fraction 0.1278 scaled to a true percent
    assert top["pct_of_aum"] == 12.78
    # market value Lakhs -> crore
    assert top["market_value_cr"] == round(115593.25 / 100, 2)


def _sbi_style_sheet():
    """SBI shape: AMC name line to skip, 'Scheme Name :' label, '% to AUM' header,
    and the statement date as a real date cell beside an 'AS ON' label."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["SBI Mutual Fund", "133"])
    ws.append(["SCHEME NAME :", "SBI Banking And Financial Services Fund"])
    ws.append(["PORTFOLIO STATEMENT AS ON", datetime.datetime(2026, 6, 30)])
    ws.append([None])
    ws.append([None, None, "Name of the Instrument / Issuer", "ISIN",
               "Rating / Industry^", "Quantity", "Market value (Rs. in Lakhs)", "% to AUM"])
    ws.append([None, None, "EQUITY & EQUITY RELATED"])
    ws.append([None, None, "HDFC Bank Ltd.", "INE040A01034", "Banks", 100, 18150.0, 18.15])
    ws.append([None, None, "ICICI Bank Ltd.", "INE090A01021", "Banks", 90, 13400.0, 13.40])
    return ws


def test_parses_sbi_style_layout():
    name, as_of, rows = base.parse_standard_sheet(_sbi_style_sheet())
    # takes the "Scheme Name :" value, not the "SBI Mutual Fund" AMC line
    assert name == "SBI Banking And Financial Services Fund"
    assert as_of == datetime.date(2026, 6, 30)   # from the date cell
    assert len(rows) == 2                          # '% to AUM' column recognised
    assert rows[0]["instrument_name"] == "HDFC Bank Ltd." and rows[0]["pct_of_aum"] == 18.15


def _kotak_style_sheet():
    """Kotak shape: title 'Portfolio of <fund> as on <date>', header 'Name of
    Instrument' (no 'the') in col A but names in col C beside the ISIN (merged-cell
    offset)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append([None, None, "Portfolio of Kotak Technology Fund as on 30-Jun-2026"])
    ws.append(["Name of Instrument", None, None, "ISIN Code", "Industry",
               "Yield", "Quantity", "Market Value (Rs.in Lacs)", "% to Net Assets"])
    ws.append(["Equity & Equity related"])
    ws.append([None, " ", "BHARTI AIRTEL LTD.", "INE397D01024", "Telecom - Services",
               None, 1000, 1835.0, 18.35])
    ws.append([None, " ", "Infosys Ltd.", "INE009A01021", "IT - Software",
               None, 900, 1200.0, 12.00])
    return ws


def test_parses_kotak_style_layout():
    name, as_of, rows = base.parse_standard_sheet(_kotak_style_sheet())
    # 'Portfolio of' prefix and ' as on <date>' suffix stripped
    assert name == "Kotak Technology Fund"
    assert as_of == datetime.date(2026, 6, 30)
    # names recovered from the column beside the ISIN despite the header offset
    assert len(rows) == 2
    assert rows[0]["instrument_name"] == "BHARTI AIRTEL LTD."
    assert rows[0]["isin"] == "INE397D01024" and rows[0]["pct_of_aum"] == 18.35


def test_instrument_name_footnotes_stripped():
    # Franklin appends markers like "$$ ~~" to names
    from app.ingestion.holdings.base import _clean_name
    assert _clean_name("HDFC Bank Ltd $$ ~~") == "HDFC Bank Ltd"
    assert _clean_name("Infosys Ltd #") == "Infosys Ltd"
    assert _clean_name("Reliance Industries Ltd") == "Reliance Industries Ltd"


def test_date_wordings():
    from app.ingestion.holdings.base import DATE_RE
    from dateutil import parser as dp
    for text in ["as on June 30, 2026", "as on June 30,2026",
                 "as on 30-Jun-2026", "as at 30 June 2026"]:
        m = DATE_RE.search(text)
        assert m, text
        assert dp.parse(m.group(1), dayfirst=True).date() == datetime.date(2026, 6, 30)
